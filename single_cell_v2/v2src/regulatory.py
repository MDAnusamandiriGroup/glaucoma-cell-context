import bisect
import json
import numpy as np
import pandas as pd
import openpyxl
from .cache import save_json
from .replication import TYPE_MAP


SUBTYPE_MAP = {k.replace(' ', '.').replace('-', '.'): v for k, v in TYPE_MAP.items()}
SUBTYPE_MAP['OFF.cone.bipolar'] = 'Bipolar'
SUBTYPE_MAP['ON.cone.bipolar'] = 'Bipolar'


def tokens(value):
    if pd.isna(value):
        return set()
    return {x.strip() for x in str(value).split(',') if x.strip()}


def canonical_types(value):
    original = tokens(value)
    unknown = original - set(SUBTYPE_MAP)
    if unknown:
        raise ValueError('Unrecognized published ATAC types: ' + str(unknown))
    return {SUBTYPE_MAP[t] for t in original}


def run(root, stages, directory, config):
    raw = root / 'data/raw'
    candidates = pd.read_csv(stages['expanded'] / 'candidate_mapping.csv')
    nominated = set(candidates['gene'])
    legacy = pd.read_csv(root / '../results/analysis_173e90345ea951f1/tables/colocalization_records.csv')
    pairs = legacy.loc[legacy['baseline_pass'] & legacy['gene'].isin(nominated), ['locus', 'rsid', 'gene_id', 'gene']].drop_duplicates()
    if pairs['gene'].nunique() != 228 or pairs['locus'].nunique() != 76:
        raise ValueError('Frozen genetic pair set changed')
    variants = json.loads((raw / 'ensembl_index_variants_GRCh38.json').read_text())
    coordinate_rows, coordinate_lookup = [], {}
    for rsid, group in pairs.groupby('rsid'):
        chromosome = str(group['locus'].iloc[0].split(':')[0])
        record = variants.get(rsid, {})
        mappings = [m for m in record.get('mappings', []) if m.get('assembly_name') == 'GRCh38'
                    and 'chr' + m.get('seq_region_name', '') == chromosome and m.get('start') == m.get('end')]
        unique = {(m['seq_region_name'], m['start']) for m in mappings}
        if len(unique) == 1:
            _, position = next(iter(unique))
            original_position = int(group['locus'].iloc[0].split(':')[1])
            if position != original_position:
                raise ValueError('Index variant coordinate disagrees with GRCh38 snapshot')
            coordinate_lookup[rsid] = (chromosome, position - 1)
            status = 'unique_GRCh38_position'
        else:
            position = None; status = 'not_returned_or_ambiguous'
        coordinate_rows.append({'rsid': rsid, 'chromosome': chromosome, 'GRCh38_position_1based': position,
            'status': status, 'returned_canonical_name': record.get('name', '')})
    pd.DataFrame(coordinate_rows).to_csv(directory / 'index_variant_coordinate_audit.csv', index=False)
    by_chromosome = {}
    for rsid, (chrom, position) in coordinate_lookup.items():
        by_chromosome.setdefault(chrom, []).append((position, rsid))
    for chrom in by_chromosome:
        by_chromosome[chrom].sort()

    def overlaps(chrom, start, end):
        values = by_chromosome.get(str(chrom), [])
        left = bisect.bisect_left(values, (int(start), ''))
        right = bisect.bisect_left(values, (int(end), ''))
        return values[left:right]

    peak_rows = []
    workbook = openpyxl.load_workbook(raw / 'wang_mmc3.xlsx', read_only=True, data_only=True)
    for sheet in workbook.sheetnames:
        if sheet in ['Overview', 'Union']:
            continue
        cell_type = SUBTYPE_MAP[sheet]
        for row in workbook[sheet].iter_rows(values_only=True):
            if row[0] is None:
                continue
            chrom, start, end = row[:3]
            for _, rsid in overlaps(chrom, start, end):
                peak_rows.append({'rsid': rsid, 'cell_type': cell_type, 'published_subtype': sheet,
                    'chromosome': chrom, 'start_0based': int(start), 'end_exclusive': int(end), 'source': 'Wang_DataS2'})
        print('ATAC intersections complete: ' + sheet, flush=True)
    workbook.close()
    peak_frame = pd.DataFrame(peak_rows, columns=['rsid', 'cell_type', 'published_subtype', 'chromosome', 'start_0based', 'end_exclusive', 'source'])
    peak_frame.to_csv(directory / 'index_variant_ATAC_peak_overlaps.csv', index=False)
    peak_types = peak_frame.groupby('rsid')['cell_type'].agg(lambda x: set(x)).to_dict()
    marker_rows = []
    workbook = openpyxl.load_workbook(raw / 'wang_mmc4.xlsx', read_only=True, data_only=True)
    for sheet in workbook.sheetnames:
        if sheet == 'Overview':
            continue
        for row in workbook[sheet].iter_rows(values_only=True):
            if row[0] is None:
                continue
            for gene in tokens(row[3]) & nominated:
                marker_rows.append({'gene': gene, 'cell_type': SUBTYPE_MAP[sheet], 'published_subtype': sheet,
                    'chromosome': row[0], 'start_0based': int(row[1]), 'end_exclusive': int(row[2]),
                    'evidence': 'Nearest-gene annotation only; does not establish a functional enhancer-gene link'})
    workbook.close()
    marker_frame = pd.DataFrame(marker_rows)
    marker_frame.to_csv(directory / 'candidate_nearest_gene_marker_peaks.csv', index=False)
    marker_frame.groupby(['gene', 'cell_type'], observed=True).size().rename('n_nearest_gene_marker_peaks').reset_index().to_csv(directory / 'candidate_nearest_gene_marker_counts.csv', index=False)
    published = pd.read_excel(raw / 'wang_mmc7.xlsx', sheet_name='Glaucoma')
    if published['ID'].duplicated().any():
        raise ValueError('Duplicate glaucoma variant annotations')
    direct = published.loc[published['ID'].isin(pairs['rsid'])].copy()
    direct.to_csv(directory / 'published_glaucoma_index_variant_annotations.csv', index=False)
    published_map = published.set_index('ID')
    original = pd.read_csv(stages['expanded'] / 'expanded_robustness.csv').set_index('gene')
    pair_rows, validations = [], []
    for record in pairs.to_dict('records'):
        rsid, gene = record['rsid'], record['gene']
        tested = rsid in published_map.index
        source = published_map.loc[rsid] if tested else None
        actual_types = peak_types.get(rsid, set())
        if tested:
            source_types = canonical_types(source['Accessible cell types'])
            validations.append({'rsid': rsid, 'computed_peak_types': ';'.join(sorted(actual_types)),
                'source_peak_types': ';'.join(sorted(source_types)),
                'types_exactly_agree': actual_types == source_types})
            if rsid in coordinate_lookup:
                chrom, pos = coordinate_lookup[rsid]
                if str(source['Chr']) != chrom or int(source['Pos']) != pos + 1:
                    raise ValueError('Published SNP coordinate mismatch')
        flags = {}
        for field, label in [('Predicted target gene(s)', 'published_predicted_target_match'),
            ('Co-accessible promoter(s)', 'published_coaccess_promoter_match'),
            ('Linked gene(s)', 'published_HiChIP_target_match'),
            ('Significantly associated gene(s)', 'published_retina_eQTL_target_match')]:
            flags[label] = gene in tokens(source[field]) if tested else False
        primary_type = original.loc[gene, 'primary_top_type']
        extended_type = original.loc[gene, 'extended_top_type']
        pair_rows.append({**record, 'in_Wang_glaucoma_table': tested,
            'gene_link_assessment_status': 'source_assessed' if tested else 'not_in_source_table',
            'computed_index_ATAC_types': ';'.join(sorted(actual_types)),
            'primary_RNA_type_in_index_ATAC': primary_type in actual_types if isinstance(primary_type, str) else False,
            'extended_RNA_type_in_index_ATAC': extended_type in actual_types if isinstance(extended_type, str) else False,
            **flags, 'any_published_gene_link': any(flags.values()),
            'source_competing_HiChIP_genes': str(source['Linked gene(s)']) if tested and pd.notna(source['Linked gene(s)']) else '',
            'source_retina_eQTL_genes': str(source['Significantly associated gene(s)']) if tested and pd.notna(source['Significantly associated gene(s)']) else ''})
    pd.DataFrame(pair_rows).to_csv(directory / 'variant_candidate_regulatory_evidence.csv', index=False)
    pd.DataFrame(validations).drop_duplicates().to_csv(directory / 'ATAC_source_reproduction_audit.csv', index=False)
    # Import all matching source predictions, including non-high-effect calls.
    scores = pd.read_excel(raw / 'wang_mmc9.xlsx', sheet_name='Disease-associated SNPs')
    scores.loc[scores['rs'].isin(pairs['rsid'])].to_csv(directory / 'published_index_BPNet_scores.csv', index=False)
    # Additional recent atlas: use identified intervals only, not unpaired row-order links.
    yuan = pd.read_csv(raw / 'yuan_table_S6.tsv', sep='\t', skiprows=1)
    yuan_types = {'All_amacrine': 'Amacrine', 'Astrocyte': 'Astrocyte', 'Cones': 'Cone',
        'GABA_amacrine': 'Amacrine', 'Gly_amacrine': 'Amacrine', 'Horizontal': 'Horizontal',
        'Muller_glia': 'Muller glia', 'OFF_cone_bipolar': 'Bipolar', 'ON_cone_bipolar': 'Bipolar',
        'RGC': 'RGC', 'Rod_bipolar': 'Bipolar', 'Rods': 'Rod'}
    if set(yuan.columns) != {'peak', 'p_value', *yuan_types}:
        raise ValueError('Unexpected Yuan retina/macula interval schema')
    recent_rows = []
    for row in yuan.to_dict('records'):
        chrom, start, stop = row['peak'].rsplit('-', 2)
        # Source IDs encode 500 bp spans with inclusive stop; use half-open stop+1.
        hits = overlaps(chrom, int(start), int(stop) + 1)
        if not hits:
            continue
        maximum = max(float(row[t]) for t in yuan_types)
        leading = [t for t in yuan_types if np.isclose(float(row[t]), maximum, rtol=0, atol=1e-10)]
        for _, rsid in hits:
            recent_rows.append({'rsid': rsid, 'peak': row['peak'],
                'source_relative_top_subtypes': ';'.join(leading),
                'canonical_top_types': ';'.join(sorted({yuan_types[t] for t in leading})),
                'source_nominal_p_value': float(row['p_value']), 'microglia_assessed_in_this_table': False,
                'scope': 'Intersection with source-listed cCREs; no new gene link or significance claim'})
    pd.DataFrame(recent_rows, columns=['rsid', 'peak', 'source_relative_top_subtypes', 'canonical_top_types',
        'source_nominal_p_value', 'microglia_assessed_in_this_table', 'scope']).to_csv(directory / 'Yuan2026_index_cCRE_overlaps.csv', index=False)
    result = pd.DataFrame(pair_rows)
    save_json(directory / 'summary.json', {'index_variants': int(pairs['rsid'].nunique()),
        'coordinate_mapped_variants': len(coordinate_lookup), 'locus_gene_pairs': len(pairs),
        'index_variants_in_Wang_glaucoma_table': int(direct['ID'].nunique()),
        'index_variants_overlapping_reproducible_ATAC': int(peak_frame['rsid'].nunique()),
        'candidate_genes_with_any_published_gene_link': int(result.loc[result['any_published_gene_link'], 'gene'].nunique()),
        'nearest_gene_marker_annotated_candidates': int(marker_frame['gene'].nunique()),
        'Yuan2026_variants_overlapping_source_listed_cCRE': len({r['rsid'] for r in recent_rows}),
        'ATAC_computed_vs_published_agreement': int(pd.DataFrame(validations).drop_duplicates()['types_exactly_agree'].sum()),
        'scope': 'Exact-index-variant intersections and separate imported regulatory evidence; no fresh fine-mapping, colocalization, MR, model training or experimentally validated causal targets',
        'genetic_independence': 'Hamel and Wang use overlapping published retinal eQTL resources; this layer is not independent genetic replication',
        'ABC_pair_status': 'Yuan separate peak/gene matrices not joined by row order; explicit enhancer-gene edge identifiers needed for that analysis'})
